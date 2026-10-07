"""Unit tests of the ingestion building blocks: parsers, chunkers, connectors, scheduler, registry."""

from __future__ import annotations

import io
import json
from datetime import datetime

import httpx
import pytest

from aiready.core.errors import ConfigError, PermissionDenied, ValidationFailed
from aiready.core.plugins import create
from aiready.core.types import Document
from aiready.ingestion.chunking import FixedChunker, HeadingChunker, TableRowsChunker
from aiready.ingestion.connectors import assert_read_only, parse_time
from aiready.ingestion.documents import decode_document, encode_document
from aiready.ingestion.enrich import chunk_id
from aiready.ingestion.parsers import Block, ParserSelector
from aiready.ingestion.registry import DocumentRegistry
from aiready.ingestion.scheduler import Schedule

# --------------------------------------------------------------------------- parsers


def parse(mime: str, content, filename: str | None = None):
    return ParserSelector().get(mime, filename).parse(content, mime_type=mime, filename=filename)


def test_markdown_headings_and_title():
    p = parse("text/markdown", "# Guide\n\nIntro.\n\n## Part A\n\nText A.\n\nSetext\n------\n\nText B.\n")
    assert p.title == "Guide"
    assert [(b.type, b.text) for b in p.blocks if b.type == "heading"] == [("heading", "Guide"), ("heading", "Part A"),
                                                                           ("heading", "Setext")]


def test_html_strips_navigation_and_keeps_tables():
    html = """<html><head><title>T</title><script>x()</script></head><body>
    <div class="site-nav">Home | Products</div><nav>Menu</nav>
    <h1>Rates</h1><p>Current rates.</p>
    <table><tr><th>Product</th><th>Rate</th></tr><tr><td>Home</td><td>1.2</td></tr></table>
    <footer>Copyright</footer></body></html>"""
    p = parse("text/html", html)
    text = p.text
    assert "Menu" not in text and "Products" not in text and "Copyright" not in text and "x()" not in text
    assert any(b.type == "row" and b.text == "Product: Home; Rate: 1.2" for b in p.blocks)
    assert p.title == "T"


def test_csv_json_rows():
    p = parse("text/csv", "a;b\n1;2\n3;4\n")
    assert [b.text for b in p.blocks] == ["a: 1; b: 2", "a: 3; b: 4"] and p.blocks[0].header == "a | b"
    j = parse("application/json", json.dumps([{"id": 1, "x": {"y": "z"}}]))
    assert j.blocks[0].text == "id: 1; x.y: z"
    with pytest.raises(ValidationFailed):
        parse("application/json", "{not json")


def test_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Handbook", level=1)
    d.add_paragraph("First paragraph.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "k", "v", "a", "1"
    buf = io.BytesIO()
    d.save(buf)
    p = parse("application/vnd.openxmlformats-officedocument.wordprocessingml.document", buf.getvalue(), "x.docx")
    assert p.blocks[0].type == "heading" and p.blocks[0].text == "Handbook"
    assert any(b.type == "row" and b.text == "k: a; v: 1" for b in p.blocks)


def test_pdf_scanned_pages_and_corrupt_pdf():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    p = parse("application/pdf", buf.getvalue())
    assert p.pages == 1 and p.metadata["scanned_pages"] == [1] and p.blocks == []
    with pytest.raises(ValidationFailed):
        parse("application/pdf", b"%PDF-garbage")


def test_parser_selection_overrides_and_unknown():
    sel = ParserSelector({".txt": "markdown"})
    assert sel.get("text/plain", "a.txt").name == "markdown"
    assert sel.get("text/x-unknown", "a.weird").name == "text"
    with pytest.raises(ValidationFailed):
        sel.get("application/octet-stream", "a.bin")


# --------------------------------------------------------------------------- chunkers


def test_heading_chunker_sections_overlap_and_tokens():
    sentences = " ".join(f"Sentence number {i} explains rule {i}." for i in range(60))
    blocks = [Block(type="heading", text="Doc", level=1), Block(type="heading", text="Rules", level=2),
              Block(text=sentences), Block(type="heading", text="Other", level=2), Block(text="Short other text here.")]
    drafts = HeadingChunker(max_tokens=60, overlap_tokens=15, min_tokens=5).chunk(blocks)
    rules = [d for d in drafts if d.section == "Doc > Rules"]
    assert len(rules) > 2 and drafts[-1].section == "Doc > Other"
    from aiready.core.tokens import count_tokens

    assert all(count_tokens(d.text) <= 60 + 5 for d in rules)  # + heading line
    # overlap: the last sentence of a window starts the next one
    first_tail = rules[0].text.split(". ")[-1].strip(". ")
    assert first_tail in rules[1].text


def test_fixed_and_table_rows_chunkers():
    blocks = [Block(text=" ".join(f"word{i}." for i in range(300)))]
    fixed = FixedChunker(max_tokens=50, overlap_tokens=10).chunk(blocks)
    assert len(fixed) > 3
    rows = [Block(type="row", text=f"id: {i}; v: x", header="id | v") for i in range(45)]
    t = TableRowsChunker(rows_per_chunk=20).chunk([Block(type="heading", text="Sheet", level=1)] + rows)
    assert len(t) == 3 and all(d.text.startswith("Columns: id | v") for d in t) and t[0].section == "Sheet"
    with pytest.raises(ValueError):
        HeadingChunker(max_tokens=10, overlap_tokens=10)


def test_chunk_ids_are_deterministic():
    assert chunk_id("d", 0, "x") == chunk_id("d", 0, "x") != chunk_id("d", 0, "y")
    assert chunk_id("d", 0, "x").startswith(chunk_id("d", 5, "z")[:17])


# --------------------------------------------------------------------------- documents / registry


def test_document_encoding_roundtrip_binary():
    doc = Document(id="a", source="s", content=b"\x00\xffbinary", checksum="c")
    assert decode_document(json.loads(json.dumps(encode_document(doc)))).content == b"\x00\xffbinary"


def test_registry_classification():
    reg = DocumentRegistry()
    doc = Document(id="d1", source="s", content="x", checksum="c1")
    assert reg.classify(doc) == "new"
    reg.mark_indexed(doc, run_id="r1", chunk_ids=["c"], collection="k")
    assert reg.classify(doc) == "unchanged"
    assert reg.classify(doc.model_copy(update={"checksum": "c2"})) == "changed"
    assert reg.classify(doc.model_copy(update={"deleted": True})) == "deleted"
    reg.mark("d1", "s", "failed", error="boom")
    assert reg.classify(doc) == "changed"  # failures are retried
    assert [e.doc_id for e in reg.missing("s", set())] == ["d1"]


# --------------------------------------------------------------------------- scheduler


def test_schedules():
    assert Schedule("every 15m").next_after(1000) == 1000 + 900
    assert Schedule("interval:2h").interval == 7200
    base = datetime(2026, 1, 5, 10, 7).timestamp()  # a Monday
    nxt = datetime.fromtimestamp(Schedule("*/15 * * * *").next_after(base))
    assert (nxt.hour, nxt.minute) == (10, 15)
    nxt = datetime.fromtimestamp(Schedule("@daily").next_after(base))
    assert (nxt.day, nxt.hour, nxt.minute) == (6, 0, 0)
    nxt = datetime.fromtimestamp(Schedule("30 6 * * 6").next_after(base))  # Saturday 06:30
    assert nxt.weekday() == 5 and (nxt.hour, nxt.minute) == (6, 30)
    assert Schedule("every 1m").is_due(None, 0) and not Schedule("every 1h").is_due(1000, 1500)
    with pytest.raises(ValueError):
        Schedule("whenever")


async def test_scheduler_runs_due_and_triggered_sources(services, docs_dir, tmp_path):
    from aiready.ingestion.config import IngestionConfig
    from aiready.ingestion.pipeline import IngestionPipeline
    from aiready.ingestion.scheduler import Scheduler

    cfg = IngestionConfig.model_validate({"sources": [{"name": "docs", "schedule": "every 1h", "collection": "k",
                                                      "connector": {"type": "filesystem", "root": str(docs_dir)}}],
                                          "landing": {"path": str(tmp_path / "l")}})
    pipe = IngestionPipeline(cfg)
    sched = Scheduler(pipe)
    assert await sched.tick_once() == ["docs"]
    assert await sched.tick_once() == []  # not due again for an hour
    await sched.trigger("docs", "webhook")
    assert await sched.tick_once() == ["docs"]
    assert pipe.registry.runs("docs")[0].trigger == "webhook"


# --------------------------------------------------------------------------- connectors


def test_read_only_sql_guard():
    assert_read_only("SELECT * FROM claims WHERE id = 1")
    for bad in ("DELETE FROM claims", "SELECT 1; DROP TABLE claims", "UPDATE claims SET x=1"):
        with pytest.raises(PermissionDenied):
            assert_read_only(bad)


async def test_sql_connector_rows_with_row_level_tags(tmp_path):
    from examples.legacy_app.db import init_db  # repo root is put on sys.path by conftest

    db = tmp_path / "replica.db"
    init_db(db).close()
    conn = create("connector", {"type": "sql", "url": f"sqlite:///{db}", "tables": [{
        "table": "claims", "title_template": "Claim {claim_number}", "owner_column": "handler",
        "content_template": "Claim {claim_number}: {description}", "access_tags": ["fraud"],
        "access_tag_columns": [{"column": "region", "template": "region:{value}"}],
        "access_rules": [{"when": {"status": "closed"}, "tags": ["service-team"]}]}]}, source="legacy-db")
    docs = {d.id: d async for d in conn.documents()}
    assert set(docs) == {"legacy-db:claims:1", "legacy-db:claims:2", "legacy-db:claims:3"}
    assert docs["legacy-db:claims:1"].access_tags == ["fraud", "region:north"]
    assert docs["legacy-db:claims:3"].access_tags == ["fraud", "region:north", "service-team"]
    assert docs["legacy-db:claims:2"].owner == "bob" and "Rear-end collision" in docs["legacy-db:claims:2"].content
    with pytest.raises(ConfigError):
        create("connector", {"type": "sql", "url": "sqlite:///x.db", "primary_url": "sqlite:///x.db", "tables": []},
               source="s")
    with pytest.raises(PermissionDenied):
        create("connector", {"type": "sql", "url": "sqlite:///x.db",
                             "tables": [{"table": "t", "query": "DELETE FROM t"}]}, source="s")


async def test_http_json_connector_paginates():
    pages = {1: {"items": [{"id": 1, "title": "a", "region": "north"}], "next_page": 2},
             2: {"items": [{"id": 2, "title": "b", "region": "south"}], "next_page": None}}
    seen_auth = []

    def handler(req: httpx.Request) -> httpx.Response:
        seen_auth.append(req.headers.get("authorization"))
        return httpx.Response(200, json=pages[int(req.url.params["page"])])

    conn = create("connector", {"type": "http_json", "base_url": "http://app", "path": "/api/items",
                                "pagination": {"type": "page", "next_path": "next_page"}, "title_template": "Item {title}",
                                "access_tag_fields": [{"field": "region", "template": "region:{value}"}],
                                "bearer_token": "tok", "transport": httpx.MockTransport(handler)}, source="api")
    docs = [d async for d in conn.documents()]
    assert [d.id for d in docs] == ["api:1", "api:2"] and docs[1].access_tags == ["region:south"]
    assert docs[0].title == "Item a" and seen_auth == ["Bearer tok", "Bearer tok"]


async def test_confluence_space_permissions_become_tags():
    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path == "/rest/api/space/CLM":
            return httpx.Response(200, json={"permissions": [
                {"operation": {"operation": "read", "targetType": "space"},
                 "subjects": {"group": {"results": [{"name": "claims-team"}]}, "user": {"results": [{"username": "eve"}]}}},
                {"operation": {"operation": "administer", "targetType": "space"},
                 "subjects": {"group": {"results": [{"name": "admins"}]}}}]})
        return httpx.Response(200, json={"results": [{
            "id": "42", "title": "Claims FAQ", "body": {"storage": {"value": "<h1>FAQ</h1><p>Answer.</p>"}},
            "version": {"when": "2026-01-01T10:00:00.000Z", "number": 3}, "_links": {"webui": "/x/42"}}],
            "_links": {}})

    conn = create("connector", {"type": "confluence", "base_url": "http://wiki", "spaces": ["CLM"],
                                "transport": httpx.MockTransport(handler)}, source="wiki")
    docs = [d async for d in conn.documents()]
    assert docs[0].access_tags == ["claims-team", "user:eve"] and docs[0].mime_type == "text/html"
    assert docs[0].updated_at == parse_time("2026-01-01T10:00:00Z") and docs[0].uri == "http://wiki/x/42"


async def test_jira_issues_with_project_and_security_level_access():
    issues = [{"key": "CLM-1", "fields": {"summary": "Broken pipe", "description": "Water", "project": {"key": "CLM"},
                                          "updated": "2026-01-01T10:00:00.000+0100", "status": {"name": "Open"},
                                          "comment": {"comments": [{"author": {"displayName": "Ann"}, "body": "ok"}]}}},
              {"key": "CLM-2", "fields": {"summary": "Suspicious", "project": {"key": "CLM"},
                                          "security": {"name": "Fraud"}}},
              {"key": "OPS-1", "fields": {"summary": "Unmapped", "project": {"key": "OPS"}}}]

    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"issues": issues, "total": 3})

    conn = create("connector", {"type": "jira", "base_url": "http://jira", "jql": "project = CLM",
                                "project_access": {"CLM": ["claims-team"]}, "security_level_access": {"Fraud": ["fraud"]},
                                "transport": httpx.MockTransport(handler)}, source="tickets")
    docs = {d.id: d async for d in conn.documents()}
    assert set(docs) == {"tickets:CLM-1", "tickets:CLM-2"}  # unmapped project fails closed
    assert docs["tickets:CLM-1"].access_tags == ["claims-team"] and "Ann: ok" in docs["tickets:CLM-1"].content
    assert docs["tickets:CLM-2"].access_tags == ["fraud"]
