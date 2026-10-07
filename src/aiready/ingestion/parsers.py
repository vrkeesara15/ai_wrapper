"""Parsers (plugin kind ``parser``): text, tables and scanned pages turned into clean text blocks.

Every parser returns a ``Parsed`` object made of ordered *blocks* (heading / text / row) rather than one
string, because the chunkers need the structure: headings to split along sections, page numbers for
citations, table rows to keep each row whole with its header.

Built-ins: ``text``, ``markdown``, ``html`` (BeautifulSoup; navigation and boilerplate stripped), ``pdf``
(pypdf, page-marked; pages without a text layer go to the OCR hook, plugin kind ``ocr``, built-in
``tesseract`` imported lazily), ``docx`` (python-docx, heading styles and tables), ``csv`` and ``xlsx``
(rows -> "column: value" text rows; openpyxl imported lazily), ``json``.

Selecting a parser: per-source overrides by mime type or extension, then the built-in mime map.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from pathlib import PurePosixPath
from typing import Any, Literal, Protocol

from pydantic import BaseModel, Field

from aiready.core.errors import ValidationFailed
from aiready.core.plugins import create, register

log = logging.getLogger(__name__)


class Block(BaseModel):
    type: Literal["heading", "text", "row"] = "text"
    text: str
    level: int = 0  # heading level (1 = top)
    page: int | None = None
    header: str | None = None  # for rows: the column header line, repeated in every chunk of rows
    metadata: dict[str, Any] = Field(default_factory=dict)


class Parsed(BaseModel):
    blocks: list[Block] = Field(default_factory=list)
    title: str | None = None
    pages: int | None = None
    parser: str = ""
    metadata: dict[str, Any] = Field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


class Parser(Protocol):
    name: str

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed: ...


def _as_text(content: bytes | str) -> str:
    if isinstance(content, str):
        return content
    for enc in ("utf-8-sig", "utf-16") if content[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _paragraphs(text: str, page: int | None = None) -> list[Block]:
    return [Block(text=p.strip(), page=page) for p in re.split(r"\n\s*\n", text) if p.strip()]


# --------------------------------------------------------------------------- plain / markdown


@register("parser", "text")
class TextParser:
    name = "text"

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        text = _as_text(content)
        if "\x00" in text:
            raise ValidationFailed("binary content cannot be parsed as text")
        return Parsed(blocks=_paragraphs(text), parser=self.name)


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_SETEXT = re.compile(r"^(=+|-+)\s*$")


@register("parser", "markdown")
class MarkdownParser:
    """ATX (``#``) and setext headings become heading blocks; fenced code is kept verbatim."""

    name = "markdown"

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        lines = _as_text(content).splitlines()
        blocks: list[Block] = []
        buf: list[str] = []
        in_fence = False
        title: str | None = None

        def flush() -> None:
            if buf:
                blocks.extend(_paragraphs("\n".join(buf)))
                buf.clear()

        for i, line in enumerate(lines):
            if line.strip().startswith("```"):
                in_fence = not in_fence
                buf.append(line)
                continue
            if not in_fence:
                m = _MD_HEADING.match(line)
                if m:
                    flush()
                    level, text = len(m.group(1)), m.group(2)
                    blocks.append(Block(type="heading", text=text, level=level))
                    title = title or (text if level == 1 else None)
                    continue
                if _SETEXT.match(line) and buf and buf[-1].strip() and i > 0:
                    heading = buf.pop()
                    flush()
                    level = 1 if line.strip().startswith("=") else 2
                    blocks.append(Block(type="heading", text=heading.strip(), level=level))
                    title = title or (heading.strip() if level == 1 else None)
                    continue
            buf.append(line)
        flush()
        return Parsed(blocks=blocks, title=title, parser=self.name)


# --------------------------------------------------------------------------- html


@register("parser", "html")
class HtmlParser:
    """Strips navigation, headers, footers, scripts and forms, keeps headings, paragraphs, lists and tables."""

    name = "html"
    STRIP = ["script", "style", "noscript", "nav", "header", "footer", "aside", "form", "iframe", "svg", "button"]
    BOILERPLATE_CLASSES = re.compile(
        r"\b(nav|navbar|menu|breadcrumbs?|footer|header|sidebar|cookie-banner|cookies?|banner|share|social)\b", re.IGNORECASE)

    def __init__(self, strip_tags: list[str] | None = None, keep_classes: list[str] | None = None):
        self.strip_tags = strip_tags or self.STRIP
        self.keep_classes = set(keep_classes or [])

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(_as_text(content), "html.parser")
        title = soup.title.get_text(strip=True) if soup.title else None
        for tag in soup(self.strip_tags):
            tag.decompose()
        for tag in soup.find_all(True):
            if getattr(tag, "decomposed", False):
                continue
            attrs = getattr(tag, "attrs", None) or {}
            marker = " ".join(attrs.get("class", []) or []) + " " + str(attrs.get("id", "")) + " " + str(attrs.get("role", ""))
            if marker.strip() and self.BOILERPLATE_CLASSES.search(marker) and not (set(attrs.get("class", []) or []) & self.keep_classes):
                if tag.name not in ("html", "body", "main", "article"):
                    tag.decompose()
        blocks: list[Block] = []
        body = soup.body or soup
        for el in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote", "table", "dd", "dt"]):
            if el.find_parent(["table"]) and el.name != "table":
                continue
            if el.name in ("p", "li", "dd", "dt") and el.find_parent(["li"]) and el.name != "li":
                continue
            if el.name == "table":
                blocks.extend(_table_rows(el))
                continue
            text = " ".join(el.get_text(" ", strip=True).split())
            if not text:
                continue
            if el.name.startswith("h"):
                blocks.append(Block(type="heading", text=text, level=int(el.name[1])))
            else:
                blocks.append(Block(text=("- " + text) if el.name == "li" else text))
        if not blocks:  # unstructured HTML: fall back to visible text
            blocks = _paragraphs(body.get_text("\n", strip=True))
        if title is None:
            title = next((b.text for b in blocks if b.type == "heading" and b.level == 1), None)
        return Parsed(blocks=blocks, title=title, parser=self.name)


def _table_rows(table: Any) -> list[Block]:
    rows = table.find_all("tr")
    if not rows:
        return []
    header = [c.get_text(" ", strip=True) for c in rows[0].find_all(["th", "td"])]
    has_header = bool(rows[0].find_all("th"))
    data_rows = rows[1:] if has_header else rows
    out = []
    for r in data_rows:
        cells = [c.get_text(" ", strip=True) for c in r.find_all(["td", "th"])]
        out.append(Block(type="row", text=_row_text(header if has_header else [], cells),
                         header=" | ".join(header) if has_header else None))
    return out


def _row_text(header: list[str], cells: list[str]) -> str:
    if header and len(header) == len(cells):
        return "; ".join(f"{h}: {v}" for h, v in zip(header, cells) if str(v).strip())
    return " | ".join(str(c) for c in cells)


# --------------------------------------------------------------------------- pdf (+ OCR hook)


class OcrEngine(Protocol):
    def image_to_text(self, image: bytes) -> str: ...


@register("ocr", "tesseract")
class TesseractOcr:
    """OCR for scanned pages. ``pytesseract`` and Pillow are imported only when a scanned page is met."""

    def __init__(self, lang: str = "eng"):
        self.lang = lang

    def image_to_text(self, image: bytes) -> str:
        import pytesseract  # optional dependency
        from PIL import Image  # optional dependency

        return str(pytesseract.image_to_string(Image.open(io.BytesIO(image)), lang=self.lang))


@register("parser", "pdf")
class PdfParser:
    """Text per page with page numbers kept on every block (for citations). Pages with no text layer are
    scanned: their images go to the OCR engine when one is configured, otherwise the page is reported in
    ``metadata.scanned_pages`` so operators can see what is missing from the index."""

    name = "pdf"

    def __init__(self, ocr: dict[str, Any] | str | None = None, min_chars_per_page: int = 20):
        self.ocr_spec = ocr
        self.min_chars = min_chars_per_page
        self._ocr: OcrEngine | None = None

    def _ocr_engine(self) -> OcrEngine | None:
        if self.ocr_spec and self._ocr is None:
            self._ocr = create("ocr", self.ocr_spec)
        return self._ocr

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        from pypdf import PdfReader
        from pypdf.errors import PyPdfError

        data = content.encode("latin-1") if isinstance(content, str) else content
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = list(reader.pages)
        except (PyPdfError, ValueError, OSError) as exc:
            raise ValidationFailed(f"unreadable PDF: {exc}") from exc
        blocks: list[Block] = []
        scanned: list[int] = []
        for i, page in enumerate(pages, start=1):
            text = (page.extract_text() or "").strip()
            if len(text) < self.min_chars:
                ocr_text = self._ocr_page(page)
                if ocr_text:
                    text = ocr_text
                elif not text:
                    scanned.append(i)
                    continue
            blocks.extend(_paragraphs(text, page=i))
        title = None
        try:
            title = (reader.metadata or {}).get("/Title")  # type: ignore[union-attr]
        except Exception:  # broken metadata must not fail the document
            title = None
        return Parsed(blocks=blocks, title=str(title) if title else None, pages=len(pages), parser=self.name,
                      metadata={"scanned_pages": scanned})

    def _ocr_page(self, page: Any) -> str:
        engine = self._ocr_engine()
        if engine is None:
            return ""
        texts = []
        for img in getattr(page, "images", []) or []:
            try:
                texts.append(engine.image_to_text(img.data))
            except Exception as exc:  # OCR failures degrade to "scanned page missing", never fail the doc
                log.warning("OCR failed: %s", exc)
        return "\n".join(t for t in texts if t.strip()).strip()


# --------------------------------------------------------------------------- office


@register("parser", "docx")
class DocxParser:
    name = "docx"

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        import docx  # python-docx
        from docx.opc.exceptions import PackageNotFoundError

        data = content.encode("latin-1") if isinstance(content, str) else content
        try:
            document = docx.Document(io.BytesIO(data))
        except (PackageNotFoundError, ValueError, KeyError) as exc:
            raise ValidationFailed(f"unreadable DOCX: {exc}") from exc
        blocks: list[Block] = []
        title = document.core_properties.title or None
        body = document.element.body
        para_by_el = {p._p: p for p in document.paragraphs}
        table_by_el = {t._tbl: t for t in document.tables}
        for el in body.iterchildren():
            if el in para_by_el:
                p = para_by_el[el]
                text = p.text.strip()
                if not text:
                    continue
                style = (p.style.name or "") if p.style is not None else ""
                m = re.match(r"Heading (\d)", style)
                if style == "Title":
                    title = title or text
                    blocks.append(Block(type="heading", text=text, level=1))
                elif m:
                    blocks.append(Block(type="heading", text=text, level=int(m.group(1))))
                else:
                    blocks.append(Block(text=text))
            elif el in table_by_el:
                rows = [[c.text.strip() for c in r.cells] for r in table_by_el[el].rows]
                if rows:
                    header = rows[0]
                    for r in rows[1:]:
                        blocks.append(Block(type="row", text=_row_text(header, r), header=" | ".join(header)))
        return Parsed(blocks=blocks, title=title, parser=self.name)


@register("parser", "csv")
class CsvParser:
    """Each row becomes ``column: value; ...`` so a chunk of rows reads well and every value stays labelled."""

    name = "csv"

    def __init__(self, delimiter: str | None = None, max_rows: int = 200_000):
        self.delimiter = delimiter
        self.max_rows = max_rows

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        text = _as_text(content)
        delim = self.delimiter
        if delim is None:
            try:
                delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                delim = ","
        reader = csv.reader(io.StringIO(text), delimiter=delim)
        rows = list(reader)
        if not rows:
            return Parsed(parser=self.name)
        header = [h.strip() for h in rows[0]]
        blocks = [Block(type="row", text=_row_text(header, r), header=" | ".join(header))
                  for r in rows[1:self.max_rows + 1] if any(c.strip() for c in r)]
        return Parsed(blocks=blocks, parser=self.name, metadata={"columns": header, "rows": len(blocks)})


@register("parser", "xlsx")
class XlsxParser:
    """Each sheet is read as CSV-like rows (``openpyxl`` imported lazily); the sheet name is a heading."""

    name = "xlsx"

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        import openpyxl  # optional dependency

        data = content.encode("latin-1") if isinstance(content, str) else content
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks: list[Block] = []
        for ws in wb.worksheets:
            rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if not rows:
                continue
            blocks.append(Block(type="heading", text=ws.title, level=1))
            header = rows[0]
            blocks += [Block(type="row", text=_row_text(header, r), header=" | ".join(header)) for r in rows[1:]]
        return Parsed(blocks=blocks, parser=self.name)


@register("parser", "json")
class JsonParser:
    """A list of objects -> one row block per object; any other JSON -> ``path: value`` lines."""

    name = "json"

    def parse(self, content: bytes | str, *, mime_type: str = "", filename: str | None = None) -> Parsed:
        try:
            data = json.loads(_as_text(content))
        except ValueError as exc:
            raise ValidationFailed(f"invalid JSON: {exc}") from exc
        if isinstance(data, list) and data and all(isinstance(x, dict) for x in data):
            blocks = [Block(type="row", text="; ".join(f"{k}: {v}" for k, v in _flat(x).items() if v not in (None, "")))
                      for x in data]
            return Parsed(blocks=blocks, parser=self.name)
        lines = [f"{k}: {v}" for k, v in _flat(data).items() if v not in (None, "")]
        title = data.get("title") if isinstance(data, dict) and isinstance(data.get("title"), str) else None
        return Parsed(blocks=_paragraphs("\n".join(lines)) if lines else [], title=title, parser=self.name)


def _flat(obj: Any, prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            out.update(_flat(v, f"{prefix}.{k}" if prefix else str(k)))
    elif isinstance(obj, list) and not all(isinstance(x, (str, int, float, bool)) or x is None for x in obj):
        for i, v in enumerate(obj):
            out.update(_flat(v, f"{prefix}[{i}]"))
    elif isinstance(obj, list):
        out[prefix or "value"] = ", ".join(str(x) for x in obj)
    else:
        out[prefix or "value"] = obj
    return out


# --------------------------------------------------------------------------- selection

MIME_MAP: dict[str, str] = {
    "text/plain": "text",
    "text/markdown": "markdown",
    "text/x-markdown": "markdown",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/csv": "csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/json": "json",
}
EXT_MAP: dict[str, str] = {
    ".txt": "text", ".md": "markdown", ".markdown": "markdown", ".html": "html", ".htm": "html", ".pdf": "pdf",
    ".docx": "docx", ".csv": "csv", ".tsv": "csv", ".xlsx": "xlsx", ".json": "json", ".log": "text",
}


class ParserSelector:
    """Chooses and caches parser instances. Overrides map a mime type or ``.ext`` to a parser spec."""

    def __init__(self, overrides: dict[str, dict[str, Any] | str] | None = None,
                 defaults: dict[str, dict[str, Any] | str] | None = None):
        self.overrides = overrides or {}
        self.defaults = defaults or {}
        self._cache: dict[str, Parser] = {}

    def spec_for(self, mime_type: str, filename: str | None) -> dict[str, Any] | str:
        ext = PurePosixPath(filename).suffix.lower() if filename else ""
        mime = (mime_type or "").split(";")[0].strip().lower()
        for table in (self.overrides, self.defaults):
            if mime in table:
                return table[mime]
            if ext and ext in table:
                return table[ext]
        if mime in MIME_MAP:
            return MIME_MAP[mime]
        if ext in EXT_MAP:
            return EXT_MAP[ext]
        if mime.startswith("text/"):
            return "text"
        raise ValidationFailed(f"no parser for mime type '{mime_type}' ({filename})")

    def get(self, mime_type: str, filename: str | None = None) -> Parser:
        spec = self.spec_for(mime_type, filename)
        key = json.dumps(spec, sort_keys=True, default=str)
        if key not in self._cache:
            self._cache[key] = create("parser", spec)
        return self._cache[key]
